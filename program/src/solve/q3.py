"""C 题 问题三：预报驱动的日内调整（三层结算）——主方案、result3 与图表。

正式主方案（v1.4，见 reports/Q3_对冲取舍实验.md）：
- 0:00 计划：组合预测 λ·官方f0 + (1−λ)·历史 EWMA，λ=0.7
- 6/12/18 调整：同组合口径的最新预报重优化未执行时段（口径 A：6:00 重优化 [6,24)、
  12:00 重优化 [12,24)、18:00 重优化 [18,24)，只承诺执行未来 6 h）
- 不使用场景对冲；联合残差对冲仅作为负结果扩展保留在 q3_ablation.py
- 储能：2 日预测窗口，每日只执行首日；日末 SOC 自由并跨日传递，规划窗口末端完全自由
  （不锚定终值）
- 执行：逐槽因果 free（q2.exec_segment_causal，无段末硬目标），缺口按 5 倍交易时刻电价紧急购电
- 结算：费用 = Σ[p·x_adj + 0.5·p·|x_plan − x_adj|] + 5·Σ p·e

输出：
- results/result3.xlsx（计划购电量 / 调整购电量 / 充放电量 / 紧急购电量）
- figures/Q3_策略费用对比.pdf、Q3_逐日紧急购电.pdf、Q3_调整量分布.pdf
- reports/RESULTS_REPORT.md 章节

运行（program/ 下）：uv run python -m solve.q3
"""
from __future__ import annotations

import time

import numpy as np
import openpyxl
import pandas as pd

import program as pm
from solve import consistency as cs
from solve import q2
from solve import q3_proto as qp
from solve.common import DATA_C, RESULTS_DIR, ROOT, T

LAM = cs.Q3_LAM
ADJ_LAM = cs.Q3_ADJ_LAM
USE_HEDGE = cs.Q3_USE_HEDGE
DAY_START = q2.REPORT_START      # 31（2025-02-01）
N_Q3 = q2.N_DAY - DAY_START      # 334


def record(section: str, data, note: str = "") -> None:
    """写结果报告：先删同名旧章节再追加（pm.record_result 为追加模式）。"""
    path = pm.reports_dir() / "RESULTS_REPORT.md"
    if path.exists():
        text = path.read_text(encoding="utf-8")
        marker = f"### {section}"
        pos = text.find(marker)
        if pos != -1:
            end = text.find("\n### ", pos + len(marker))
            text = text[:pos] if end == -1 else text[:pos] + text[end + 1:]
            path.write_text(text, encoding="utf-8")
    pm.record_result(section, data, note=note)


def _day_result(data, D, e_start, **kw):
    if USE_HEDGE:  # 仅用于显式复现实验；正式唯一参数源固定为 False。
        r = qp.simulate_day_rt_hedge(
            data, D, LAM, adj_lam=ADJ_LAM, n_scen=cs.N_SCEN, seed=7,
            e_start=e_start, **kw,
        )
    else:
        r = qp.simulate_day_rt(
            data, D, LAM, adj_lam=ADJ_LAM, e_start=e_start, **kw,
        )
    return r


def run_main(data):
    """主方案全年模拟，返回明细。"""
    days = list(range(DAY_START, DAY_START + N_Q3))
    out = {k: np.zeros((N_Q3, T)) for k in ("xp", "xa", "c", "d", "E", "e")}
    out["E_start"] = np.zeros(N_Q3)
    rows = []
    t0 = time.time()
    E = float(qp.E0)
    for i, D in enumerate(days):
        r = _day_result(data, D, E)
        out["xp"][i] = r["x_plan"]; out["xa"][i] = r["x_final"]
        out["c"][i] = r["c"]; out["d"][i] = r["d"]
        out["E"][i] = r["E_exec"]; out["e"][i] = r["e"]
        out["E_start"][i] = r["E_start"]
        rows.append({"日期": data["dates"][D], "计划费/元": r["plan_cost"],
                     "调整净额/元": r["adjust_net"], "紧急费/元": r["emerg"],
                     "总费用/元": r["total"], "紧急量/kWh": float(r["e"].sum()),
                     "调整量/kWh": float(np.abs(r["x_final"] - r["x_plan"]).sum()),
                     "场景最大日序号": int(r.get("scenario_max_day", -1)),
                     "场景数": int(r.get("scenario_count", 0))})
        E = float(r["E_end"])
        if (i + 1) % 50 == 0:
            print(f"  主方案 {i + 1}/{N_Q3} 天，用时 {time.time() - t0:.0f}s")
    return {k: v for k, v in out.items()}, pd.DataFrame(rows)


def run_baseline(data, tag, **kw):
    """基线模拟（λ=1 官方，无对冲），用于对照。"""
    days = list(range(DAY_START, DAY_START + N_Q3))
    xa = np.zeros((N_Q3, T)); e = np.zeros((N_Q3, T))
    tot = plan = adj = em = 0.0
    E = float(qp.E0)
    for i, D in enumerate(days):
        r = qp.simulate_day_rt(data, D, 1.0, e_start=E, **kw)
        xa[i] = r["x_final"]; e[i] = r["e"]
        tot += r["total"]; plan += r["plan_cost"]; adj += r["adjust_net"]; em += r["emerg"]
        E = float(r["E_end"])
    print(f"  {tag}: 总 {tot / 1e4:.1f} 万（计划 {plan / 1e4:.1f} + 调整 {adj / 1e4:+.1f} + 紧急 {em / 1e4:.1f}）")
    return {"total": tot, "plan": plan, "adj": adj, "emerg": em, "e": e, "xa": xa}


def write_result3(dates, price, out):
    """按附件 5 模板生成 results/result3.xlsx。"""
    wb = openpyxl.load_workbook(DATA_C / "附件5" / "result3.xlsx")
    i0 = DAY_START

    ws = wb["计划购电量"]
    for i in range(N_Q3):
        row = 2 + i
        for t in range(T):
            ws.cell(row=row, column=2 + t, value=float(out["xp"][i, t]))
        ws.cell(row=row, column=146, value=float(out["xp"][i].sum()))
        ws.cell(row=row, column=147, value=float(price @ out["xp"][i]))

    ws = wb["调整购电量"]
    for i in range(N_Q3):
        row = 2 + i
        xa, xp = out["xa"][i], out["xp"][i]
        for t in range(T):
            ws.cell(row=row, column=2 + t, value=float(xa[t]))
        ws.cell(row=row, column=146, value=float(xa.sum()))
        settle = float((price * xa + 0.5 * price * np.abs(xa - xp)).sum())
        ws.cell(row=row, column=147, value=settle)

    ws = wb["充放电量"]
    while ws.max_row > 1:
        ws.delete_rows(2)
    row = 2
    for i in range(N_Q3):
        for b in range(6):
            rr = row + b
            if b == 0:
                ws.cell(row=rr, column=1, value=pd.Timestamp(dates[i0 + i]).to_pydatetime())
            ws.cell(row=rr, column=2, value=f"{4 * b}:00-{4 * (b + 1)}:00")
            ws.cell(row=rr, column=3, value=float(out["c"][i, 24 * b:24 * (b + 1)].sum()))
            ws.cell(row=rr, column=4, value=float(out["d"][i, 24 * b:24 * (b + 1)].sum()))
            if b == 0:
                ws.cell(row=rr, column=5, value="00:00")
                ws.cell(row=rr, column=6, value=float(out["E_start"][i]))
            if b == 1:
                ws.cell(row=rr, column=5, value="24:00")
                ws.cell(row=rr, column=6, value=float(out["E"][i, -1]))
        row += 6

    ws = wb["紧急购电量"]
    while ws.max_row > 1:
        ws.delete_rows(2)
    row = 2
    for i in range(N_Q3):
        for j, (a, b, kwh) in enumerate(q2._events(out["e"][i])):
            if j == 0:
                ws.cell(row=row, column=1, value=pd.Timestamp(dates[i0 + i]).to_pydatetime())
            ws.cell(row=row, column=2, value=f"{q2._fmt_time(a * 10)}-{q2._fmt_time((b + 1) * 10)}")
            ws.cell(row=row, column=3, value=round(kwh, 4))
            row += 1

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    p = RESULTS_DIR / "result3.xlsx"
    wb.save(p)
    wb.close()
    return p


def verify_result3(path, dates, price, out):
    """回读校验：行数、端点、勾稽、紧急事件。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    chk = {"sheets": "/".join(wb.sheetnames)}

    ws = wb["计划购电量"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    chk["计划表行数"] = len(rows)
    chk["计划购电量勾稽差/kWh"] = float(abs(sum(float(r[145]) for r in rows) - out["xp"].sum()))
    chk["计划购电费勾稽差/元"] = float(abs(sum(float(r[146]) for r in rows)
                                       - sum(float(price @ out["xp"][i]) for i in range(N_Q3))))

    ws = wb["调整购电量"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    chk["调整表行数"] = len(rows)
    chk["调整购电量勾稽差/kWh"] = float(abs(sum(float(r[145]) for r in rows) - out["xa"].sum()))

    ws = wb["充放电量"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] is not None]
    chk["充放表行数"] = len(rows)
    e0 = [float(rows[6 * i][5]) for i in range(N_Q3)]
    e24 = [float(rows[6 * i + 1][5]) for i in range(N_Q3)]
    chk["跨日SOC衔接最大误差/kWh"] = float(
        max(abs(e24[i] - e0[i + 1]) for i in range(N_Q3 - 1))
    )
    chk["日末SOC非固定取值数"] = int(len(set(round(v, 3) for v in e24)))

    ws = wb["紧急购电量"]
    kw = sum(float(r[2]) for r in ws.iter_rows(min_row=2, values_only=True) if r[2] is not None)
    chk["紧急表电量勾稽差/kWh"] = float(abs(kw - out["e"].sum()))
    wb.close()
    return chk


def make_figures(daily, base_noadj, base_off, out, dates):
    import matplotlib.pyplot as plt

    # 1) 策略费用对比（三策略 × 三构成）
    names = ["无调整（官方）", "三点调整（官方）", "主方案（组合无对冲）"]
    plan = [base_noadj["plan"] / 1e4, base_off["plan"] / 1e4,
            float(sum(daily["计划费/元"])) / 1e4]
    adj = [base_noadj["adj"] / 1e4, base_off["adj"] / 1e4,
           float(sum(daily["调整净额/元"])) / 1e4]
    em = [base_noadj["emerg"] / 1e4, base_off["emerg"] / 1e4,
          float(sum(daily["紧急费/元"])) / 1e4]
    pm.bar_group(names, {"计划购电费": plan, "调整净额": adj, "紧急购电费": em},
                 ylabel="334 天费用 / 万元", save="Q3_策略费用对比")

    # 2) 逐日紧急购电量
    fig, ax = pm.line(np.arange(N_Q3), [out["e"].sum(axis=1), base_off["e"].sum(axis=1)],
                      labels=["主方案（组合无对冲）", "三点调整（官方）"],
                      xlabel="日期（2025-02-01 起）", ylabel="紧急购电量 / kWh")
    pm.save_fig(fig, "Q3_逐日紧急购电",
                data=pd.DataFrame({"日期": daily["日期"], "主方案紧急量/kWh": out["e"].sum(axis=1),
                                   "官方三点紧急量/kWh": base_off["e"].sum(axis=1)}))

    # 3) 调整量分布（主方案：|x_adj − x_plan| 的逐日分布）
    dev = np.abs(out["xa"] - out["xp"]).sum(axis=1)
    fig2, ax2 = pm.hist(dev, bins=30, xlabel="逐日调整量 |x_adj − x_plan| / kWh",
                        ylabel="天数")
    pm.save_fig(fig2, "Q3_调整量分布", data=pd.DataFrame({"逐日调整量_kWh": dev}))


def main():
    pm.init(seed=42, root=str(ROOT))
    log = pm.get_logger("q3")
    t0 = time.time()
    data = qp.load_extended()  # 含 EWMA h=5 统一权重（consistency.py）
    print("基线对照……")
    base_noadj = run_baseline(data, "无调整（官方）", adj_hours=())
    base_off = run_baseline(data, "三点调整（官方）")
    print("主方案……")
    out, daily = run_main(data)
    cost = {"plan": float(sum(daily["计划费/元"])), "adj": float(sum(daily["调整净额/元"])),
            "emerg": float(sum(daily["紧急费/元"])), "total": float(sum(daily["总费用/元"]))}
    print(f"主方案：总 {cost['total'] / 1e4:.1f} 万（计划 {cost['plan'] / 1e4:.1f} "
          f"+ 调整 {cost['adj'] / 1e4:+.1f} + 紧急 {cost['emerg'] / 1e4:.1f}），"
          f"用时 {time.time() - t0:.0f}s")

    p = write_result3(data["dates"], data["price"], out)
    chk = verify_result3(p, data["dates"], data["price"], out)
    log.info("result3.xlsx -> {}", p)

    make_figures(daily, base_noadj, base_off, out, data["dates"])
    daily.to_csv(ROOT / "code" / "outputs" / "q3_daily.csv", index=False, encoding="utf-8-sig")

    day_idx = np.arange(DAY_START, DAY_START + N_Q3)
    chk["场景池前视违规天数"] = int(
        (daily["场景最大日序号"].to_numpy() >= day_idx).sum()
    )
    chk["场景对冲是否启用"] = bool(USE_HEDGE)

    record(
        "问题三 主方案结果（2025-02-01 ~ 12-31，334 天）",
        {
            "计划购电费/万元": round(cost["plan"] / 1e4, 1),
            "调整净额/万元": round(cost["adj"] / 1e4, 1),
            "紧急购电费/万元": round(cost["emerg"] / 1e4, 1),
            "总费用/万元": round(cost["total"] / 1e4, 1),
            "对照·无调整（官方）/万元": round(base_noadj["total"] / 1e4, 1),
            "对照·三点调整（官方）/万元": round(base_off["total"] / 1e4, 1),
            "主方案相对无调整": f"{100 * (cost['total'] / base_noadj['total'] - 1):+.2f}%",
            "主方案相对官方三点": f"{100 * (cost['total'] / base_off['total'] - 1):+.2f}%",
            "发生紧急购电天数": int((daily["紧急量/kWh"] > 1e-3).sum()),
            "全程紧急购电量/kWh": float(daily["紧急量/kWh"].sum()),
            "全程调整量/kWh": float(daily["调整量/kWh"].sum()),
            **chk,
        },
        note=(
            f"统一口径（consistency.py）：负荷使用目标日前 EWMA h=5 权重预测并抬升 "
            f"κ={cs.KAPPA:g}，"
            f"0:00 光伏用组合预测（λ={LAM:g}·官方f0 + {1 - LAM:g}·历史 EWMA）"
            f"并折减 m={cs.MARGIN:g} kW；"
            "6/12/18 点用同组合口径的最新预报调整；正式方案不使用场景对冲，"
            "联合残差对冲仅保留为消融扩展；"
            "储能采用 2 日预测窗口滚动优化，当前日末 SOC 不固定并传递到下一日，"
            "规划窗口末端完全自由（不锚定终值）；"
            "逐槽因果执行（q2.exec_segment_causal，"
            "不读取未来实际值），缺口 5 倍紧急。结算 = Σ[p·x_adj + 0.5p|x_plan−x_adj|] + 5Σp·e。"
            "result3.xlsx 已生成并回读校验；图 figures/Q3_策略费用对比.pdf、Q3_逐日紧急购电.pdf、"
            "Q3_调整量分布.pdf；逐日表 code/outputs/q3_daily.csv。"
        ),
    )
    print("完成。")


if __name__ == "__main__":
    main()
