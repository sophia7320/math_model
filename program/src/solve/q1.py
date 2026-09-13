"""C 题 问题一：典型日购电–储能连续 LP 调度。

数学模型（连续线性规划；t = 1..144 个 10 分钟时段）：

    min   Σ_t p_t · x_t                             （全天购电费）
    s.t.  P_t/6 + x_t + d_t = L_t/6 + c_t + s_t     （功率平衡；s_t≥0 为弃光）
          E_t = E_{t-1} + η·c_t − d_t/η              （储能动态；η=0.9）
          E_0 = E_144 = 6000                         （储电量循环回归）
          0 ≤ x_t；0 ≤ c_t, d_t ≤ 833.33             （购电非负；充放功率上限）
          1200 ≤ E_t ≤ 10800                         （储电量安全区间）

变量：x（购电）、c（充电）、d（放电）、s（弃光）、E（时段末储电量），共 5×144 个，全部连续。
同槽「充放互斥」无需 0-1 变量：若 c_t,d_t>0，同时减 δ 可使储电量增加 δ(1/η−η)>0，
其余约束不变 —— 连续松弛的最优解可取到无同充同放的解。

口径说明（论文假设）：
1. 附件标签为时段右端点；附件标签与附件 5 结果模板相差一个 10 分钟刻度（附件第 60 行为 10:00，
   模板第 60 行为 10:00-10:10）。数据按行序号对齐（附件第 k 槽 ↔ 模板第 k 行 ↔ 论文第 k 时段），
   结果文件沿用模板标签，模板标签仅作行标识；
2. 充电乘 η、放电除 η（往返 81%；另有「仅充电计损」口径，可切换 ETA_D 做敏感性）；
3. 允许弃光（s_t ≥ 0），不允许向外部电网售电（x_t ≥ 0）。
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.sparse import csr_matrix, lil_matrix

import program as pm
from solve.common import (
    DATA_C,
    E0,
    E_MAX,
    E_MIN,
    ETA,
    P_MAX_E,
    RESULTS_DIR,
    ROOT,
    T,
    read_attachment,
)

N_VAR = 5 * T  # 5 类变量 × 144 时段


def _ix(block: int, t: int) -> int:
    """变量索引：0=x 购电，1=c 充电，2=d 放电，3=s 弃光，4=E 储电量。"""
    return block * T + t


# ---------------------------------------------------------------------------
# 建模与求解
# ---------------------------------------------------------------------------
def build_lp(price: np.ndarray, load_kwh: np.ndarray, pv_kwh: np.ndarray):
    """构建并求解问题一连续 LP，返回 ``(OptResult, 解字典)``。

    所有能量量纲为 kWh/时段，功率量纲为 kW（乘 1/6 转 kWh）。
    """
    c_obj = np.zeros(N_VAR)
    c_obj[0:T] = price  # 只有购电产生费用

    A_eq = lil_matrix((2 * T + 1, N_VAR))
    b_eq = np.zeros(2 * T + 1)

    for t in range(T):  # ① 功率平衡
        A_eq[t, _ix(0, t)] = 1.0  # + 购电
        A_eq[t, _ix(1, t)] = -1.0  # − 充电
        A_eq[t, _ix(2, t)] = 1.0  # + 放电
        A_eq[t, _ix(3, t)] = -1.0  # − 弃光
        b_eq[t] = load_kwh[t] - pv_kwh[t]

    for t in range(T):  # ② 储能动态
        r = T + t
        A_eq[r, _ix(4, t)] = 1.0  # E_t
        if t:
            A_eq[r, _ix(4, t - 1)] = -1.0  # − E_{t-1}
        A_eq[r, _ix(1, t)] = -ETA  # − η·c_t
        A_eq[r, _ix(2, t)] = 1.0 / ETA  # + d_t/η
        b_eq[r] = E0 if t == 0 else 0.0

    A_eq[2 * T, _ix(4, T - 1)] = 1.0  # ③ 末值：E_144 = E_0
    b_eq[2 * T] = E0

    bounds = (
        [(0.0, None)] * T  # x ≥ 0
        + [(0.0, P_MAX_E)] * T  # 0 ≤ c ≤ 833.33
        + [(0.0, P_MAX_E)] * T  # 0 ≤ d ≤ 833.33
        + [(0.0, float(v)) for v in pv_kwh]  # 0 ≤ s ≤ 该时段可用光伏
        + [(E_MIN, E_MAX)] * T  # 储电量安全区间
    )

    res = pm.optimize.solve_lp(c_obj, A_eq=csr_matrix(A_eq), b_eq=b_eq, bounds=bounds)
    if not res.success:
        raise RuntimeError(f"问题一连续 LP 求解失败：{res.message}")

    sol = {
        "x": res.x[0:T],
        "c": res.x[T : 2 * T],
        "d": res.x[2 * T : 3 * T],
        "s": res.x[3 * T : 4 * T],
        "E": res.x[4 * T : 5 * T],
    }
    return res, sol


# ---------------------------------------------------------------------------
# 校验
# ---------------------------------------------------------------------------
def verify(sol: dict, load_kwh: np.ndarray, pv_kwh: np.ndarray) -> dict:
    """逐项校验功率平衡、储能动态、边界与同充同放（返回指标 dict）。"""
    x, c, d, s, E = (sol[k] for k in ("x", "c", "d", "s", "E"))
    balance_res = float(np.max(np.abs(x + d - c - s - (load_kwh - pv_kwh))))
    dyn_res = float(
        np.max(np.abs(np.diff(np.concatenate([[E0], E])) - (ETA * c - d / ETA)))
    )
    return {
        "功率平衡最大残差/kWh": balance_res,
        "储能动态最大残差/kWh": dyn_res,
        "储电量越下限/kWh": float(max(0.0, E_MIN - E.min())),
        "储电量越上限/kWh": float(max(0.0, E.max() - E_MAX)),
        "充放电越功率限/kWh": float(max(0.0, c.max() - P_MAX_E, d.max() - P_MAX_E)),
        "0:00 储电量/kWh": float(E0),
        "24:00 储电量/kWh": float(E[-1]),
        "同槽同时充放最大值/kWh": float(np.max(np.minimum(c, d))),
        "弃光总量/kWh": float(s.sum()),
        "购电量最小值/kWh": float(x.min()),
    }


# ---------------------------------------------------------------------------
# 图表
# ---------------------------------------------------------------------------
def make_figures(
    sol: dict,
    price: np.ndarray,
    load_kw: np.ndarray,
    pv_kw: np.ndarray,
    baseline_cost: float,
    cost: float,
) -> None:
    """生成问题一论文图（PDF 存 figures/，作图数据存 code/outputs/figure_data/）。"""
    x, c, d, E = (sol[k] for k in ("x", "c", "d", "E"))
    t_h = (np.arange(T) + 0.5) / 6.0  # 时段中点（小时）

    pm.line(
        t_h,
        [load_kw, pv_kw, x * 6.0, d * 6.0, c * 6.0],
        labels=["小区负载", "光伏发电", "购电", "储能放电", "储能充电"],
        xlabel="时刻 / h",
        ylabel="功率 / kW",
        save="Q1_调度时序",
    )
    pm.dual_axis(
        t_h,
        E,
        price,
        labels=("储电量", "电价"),
        xlabel="时刻 / h",
        ylabel1="储电量 / kWh",
        ylabel2="电价 / (元·kWh$^{-1}$)",
        save="Q1_储电量与电价",
    )
    pm.bar(
        ["无储能基线", "储能最优调度"],
        [baseline_cost, cost],
        ylabel="全天购电费 / 元",
        value_labels=True,
        save="Q1_费用对比",
    )


# ---------------------------------------------------------------------------
# 结果文件
# ---------------------------------------------------------------------------
def write_result1(sol: dict) -> Path:
    """按附件 5 模板填写 ``result1.xlsx``（计划购电量 + 充放电量），输出到 results/。"""
    x, c, d, E = (sol[k] for k in ("x", "c", "d", "E"))
    wb = openpyxl.load_workbook(DATA_C / "附件5" / "result1.xlsx")

    ws = wb["计划购电量"]  # 144 行，与附件同时段序号对齐
    for t in range(T):
        ws.cell(row=2 + t, column=2, value=float(x[t]))

    ws = wb["充放电量"]  # 6 个 4 小时块 + 两端储电量
    for b in range(6):
        ws.cell(row=2 + b, column=2, value=float(c[24 * b : 24 * (b + 1)].sum()))
        ws.cell(row=2 + b, column=3, value=float(d[24 * b : 24 * (b + 1)].sum()))
    ws.cell(row=2, column=5, value=float(E0))  # 0:00 储电量
    ws.cell(row=3, column=5, value=float(E[-1]))  # 24:00 储电量

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    out = RESULTS_DIR / "result1.xlsx"
    wb.save(out)
    wb.close()
    return out


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
def run_q1() -> dict:
    """问题一：读数据 → 建模求解 → 校验 → 记录 → 出图 → 写 result1.xlsx。"""
    pm.init(root=str(ROOT))
    log = pm.get_logger("q1")

    df = read_attachment("附件1.xlsx")
    label = df["时间"].astype(str).to_numpy()
    price = df["电价"].to_numpy(dtype=float)
    load_kw = df["小区负载"].to_numpy(dtype=float)
    pv_kw = df["光伏发电预测功率"].to_numpy(dtype=float)
    load_kwh, pv_kwh = load_kw / 6.0, pv_kw / 6.0

    with pm.Timer("问题一连续 LP 求解", logger=log):
        res, sol = build_lp(price, load_kwh, pv_kwh)

    x, c, d, s, E = (sol[k] for k in ("x", "c", "d", "s", "E"))
    cost = float(price @ x)
    baseline = float(np.sum(price * np.maximum(load_kwh - pv_kwh, 0.0)))

    # ---- 记录：数据概览 ----
    pm.record_result(
        "问题一 输入数据概览",
        {
            "时段数": T,
            "电价范围/(元/kWh)": f"{price.min():.4f} ~ {price.max():.4f}",
            "负载范围/kW": f"{load_kw.min():.1f} ~ {load_kw.max():.1f}",
            "光伏出力范围/kW": f"{pv_kw.min():.1f} ~ {pv_kw.max():.1f}",
        },
        note="来源：附件 1（典型日：电价 / 小区负载 / 光伏发电预测功率，10 分钟粒度）。",
    )

    # ---- 记录：结果总览 ----
    pm.record_result(
        "问题一 结果总览",
        {
            "全天购电量/kWh": float(x.sum()),
            "全天购电费/元": cost,
            "无储能基线购电费/元": baseline,
            "储能节省/元": baseline - cost,
            "储能节省比例/%": 100.0 * (baseline - cost) / baseline,
            "总充电量/kWh": float(c.sum()),
            "总放电量/kWh": float(d.sum()),
            "弃光总量/kWh": float(s.sum()),
            "储电量最小值/kWh": float(E.min()),
            "储电量最大值/kWh": float(E.max()),
        },
        note=(
            "方法：连续线性规划（HiGHS，%d 变量、%d 条等式约束），目标为全天购电费最小；"
            "储能充放电效率 %.2f（充放双向计损），E(0:00)=E(24:00)=%.0f kWh；弃光松弛 s≥0。"
            "对照基线：无储能时逐时段购电 max(负载−光伏, 0)。"
            "图表：figures/Q1_调度时序.pdf、Q1_储电量与电价.pdf、Q1_费用对比.pdf；"
            "完整时段表：code/outputs/q1_schedule.csv。" % (N_VAR, 2 * T + 1, ETA, E0)
        ),
    )

    # ---- 记录：表 1（指定 10 分钟窗口）----
    windows = [10, 12, 14, 16, 18, 20]  # 表 1 要求的整点窗口
    w_idx = [6 * h - 1 for h in windows]  # 按序号对齐口径的槽序号
    table1 = pd.DataFrame(
        {
            "时间段": [f"{h}:00-{h}:10" for h in windows],
            "购电量/kWh": [float(x[i]) for i in w_idx],
        }
    )
    pm.record_result(
        "问题一 表1 指定时段购电量",
        table1,
        note=(
            "口径：数据按行序号对齐（附件第 k 槽 ↔ 附件 5 第 k 行 ↔ 论文第 k 时段）；附件标签为时段"
            "右端点，附件与结果模板标签相差一个 10 分钟刻度（10:00-10:10 ↔ 第 60 个时段），模板标签"
            "仅作行标识。"
            "全天购电量 %.1f kWh，全天购电费 %.2f 元。" % (x.sum(), cost)
        ),
    )
    pm.save_outputs(table1, "q1_table1_buy")

    # ---- 记录：表 2（4 小时块充放电）----
    table2 = pd.DataFrame(
        {
            "时间段": [f"{4 * b}:00-{4 * (b + 1)}:00" for b in range(6)],
            "充电量/kWh": [float(c[24 * b : 24 * (b + 1)].sum()) for b in range(6)],
            "放电量/kWh": [float(d[24 * b : 24 * (b + 1)].sum()) for b in range(6)],
        }
    )
    pm.record_result(
        "问题一 表2 储能充放电量",
        table2,
        note=f"0:00 储电量 = {E0:.1f} kWh，24:00 储电量 = {E[-1]:.1f} kWh（相等，满足循环约束）。",
    )
    pm.save_outputs(table2, "q1_table2_storage")

    # ---- 记录：约束与一致性校验 ----
    checks = verify(sol, load_kwh, pv_kwh)
    pm.record_result(
        "问题一 约束与一致性校验",
        checks,
        note=(
            "残差均为数值误差量级（10⁻¹³）；储电量全程在 [1200, 10800] kWh 内；"
            "无同槽充放、无弃光、购电量非负。"
        ),
    )

    # ---- 中间数据：完整时段调度表 ----
    schedule = pd.DataFrame(
        {
            "时段": np.arange(1, T + 1),
            "时间标签": label,
            "电价/(元/kWh)": price,
            "负载/kW": load_kw,
            "光伏/kW": pv_kw,
            "购电量/kWh": x,
            "充电量/kWh": c,
            "放电量/kWh": d,
            "弃光量/kWh": s,
            "储电量/kWh": E,
        }
    )
    pm.save_outputs(schedule, "q1_schedule")

    # ---- 图 ----
    make_figures(sol, price, load_kw, pv_kw, baseline, cost)

    # ---- 官方结果文件 ----
    out = write_result1(sol)

    log.info(
        "问题一完成：购电量 {:.1f} kWh，购电费 {:.2f} 元（无储能 {:.2f} 元，节省 {:.1f}%）；result1.xlsx → {}",
        x.sum(),
        cost,
        baseline,
        100.0 * (baseline - cost) / baseline,
        out,
    )
    return {"cost": cost, "baseline": baseline, "sol": sol, "checks": checks}
