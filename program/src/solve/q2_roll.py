"""C 题 问题二（滚动口径修订，2026-09-13）：2 日规划 + 1 日执行、储电量末端自由。

修订动机：原口径把"每天 0:00 与 24:00 储电量相同（=6000 kWh）"作为储能终端条件
（日循环），但题面只在问题 1 提出该要求，问题 2/3 并未要求；强制日末回 6000 会
人为锁死储能、截断跨日价值。新口径改为滚动 MPC：

- 规划窗口：每天 0:00 以"当日 + 次日"共 2 天（288 槽）为窗口求解最小购电费用 LP，
  末端不设限（free final SOC），SOC 满足 [1200, 10800]；
- 执行窗口：只执行当日（1 天），次日 0:00 用新的状态与预测重新规划；
- 储电量跨日连续：当日执行的末端储电量即次日规划初值；
- 预测：沿用口径 E 调优配置（EWMA 半衰期 5 天标定权重 + κ=1.02 负荷抬升 +
  m=50 kW 光伏折扣）；窗口次日的预测只用决策时点可得信息（无前视，见 ``forecast_two_days``）；
- 执行：逐槽因果执行（``q2.exec_segment_causal``，只读当前已实现值），当日末储能
  跟踪规划轨迹（可达性投影），缺口按 5 倍交易时刻电价紧急购电；计划购电 take-or-pay；
- 状态：仿真自 2025-01-16 0:00 起（前 15 天为预测/权重预热，无前视），初始储电量
  6000 kWh；报送期 2025-02-01 ~ 12-31（334 天）。窗口末日（12-31）的"次日"用当日
  复制填充，避免年末终值清算效应。

运行（program/ 下，先设 $env:PYTHONIOENCODING='utf-8'）：
    uv run python -m solve.q2_roll --probe              # 结构对比（日循环 / 2日锚定 / 2日自由）
    uv run python -m solve.q2_roll --result2            # 生成官方 results/result2.xlsx
    uv run python -m solve.q2_roll --result2 --mc 100   # 并做 100 年蒙特卡洛分布
"""
from __future__ import annotations

import argparse
import shutil
import time

import numpy as np
import openpyxl
import pandas as pd

import program as pm
from solve import q2
from solve.common import E0, E_MAX, E_MIN, ETA, RESULTS_DIR, ROOT, T
from solve.q2_adaptive import AdaptiveWeightModel

N_DAY = q2.N_DAY
REPORT_START = q2.REPORT_START
START_SIM = 15          # 滚动仿真起点（日序号，0 基；前 15 天为预测预热）
HL = 5.0                # EWMA 半衰期（天，调优配置）
KAPPA = 1.02            # 负荷抬升系数（调优配置）
MARGIN = 50.0           # 光伏折扣（kW，调优配置）
EPS_PLAN = AdaptiveWeightModel.EPS_PLAN
EMERG_MULT = q2.EMERG_MULT
REP = list(range(REPORT_START, N_DAY))


def record(section: str, data, note: str = "") -> None:
    """写结果报告：先删同名旧章节再追加（pm.record_result 为追加模式）。"""
    path = pm.reports_dir() / "RESULTS_REPORT.md"
    if path.exists():
        text = path.read_text(encoding="utf-8")
        marker = f"### {section}"
        pos = text.find(marker)
        if pos != -1:
            end = text.find("\n### ", pos + len(marker))
            text = text[:pos] if end == -1 else text[:pos] + text[end + 1:]
            path.write_text(text, encoding="utf-8")
    pm.record_result(section, data, note=note)


# ---------------------------------------------------------------------------
# 权重序列（EWMA 标定，扩展到仿真起点）
# ---------------------------------------------------------------------------
def seqs_ewma_ext(model: AdaptiveWeightModel, hl: float = HL,
                  start: int = START_SIM) -> dict[int, tuple[np.ndarray, np.ndarray]]:
    """逐日 EWMA 标定权重（与 q2_tune.seqs_ewma 同式，扩展到 start 日起）。

    目标日 d 的记分矩阵 M = 衰减加权 [START, d-1] 的费用表行；
    仅用 d-1 及之前的信息（无前视）。d=31 起与官方 seqs_ewma 完全一致。
    """
    decay = float(np.exp(-np.log(2.0) / hl))
    M = model.C[model.START].copy()
    out: dict[int, tuple[np.ndarray, np.ndarray]] = {}
    for d in range(model.START + 1, N_DAY):
        M = decay * M + (1.0 - decay) * model.C[d - 1]
        if d >= start:
            i, j = np.unravel_index(np.argmin(M), M.shape)
            out[d] = (model.grid[i].copy(), model.grid[j].copy())
    return out


# ---------------------------------------------------------------------------
# 2 日窗口预测（次日无前视）
# ---------------------------------------------------------------------------
def forecast_two_days(model: AdaptiveWeightModel, w, u, D: int,
                      pv_next: str = "recurse"):
    """窗口 [D, D+1] 的负荷/光伏预测（kW，未裁剪；第二日只用 ≤D-1 信息）。

    当日（D）：口径 E 三源加权 L̂(D)、P̂(D)；
    次日（D+1）：负荷用同一权重前移一天 L(D-6)/L(D-13)/典型日；
    光伏因 P(D) 未知，用递推外推 P̂(D+1)=u1·P̂(D)+u2·P(D−1)+u3·P̄（"recurse"），
    或直接沿用当日同式值（"same"，两者都无前视）。
    """
    L, P, Lt, Pt = model.L, model.P, model.L_typ, model.P_typ
    l0 = w[0] * L[D - 7] + w[1] * L[D - 14] + w[2] * Lt
    p0 = u[0] * P[D - 1] + u[1] * P[D - 2] + u[2] * Pt
    D1 = min(D + 1, N_DAY - 1)
    if D1 == D:
        return l0, p0, l0, p0
    l1 = w[0] * L[D1 - 7] + w[1] * L[D1 - 14] + w[2] * Lt
    p1 = (u[0] * p0 + u[1] * P[D - 1] + u[2] * Pt) if pv_next == "recurse" else p0.copy()
    return l0, p0, l1, p1


# ---------------------------------------------------------------------------
# 全年滚动（三种结构：日循环 / 2日锚定 / 2日自由）
# ---------------------------------------------------------------------------
def run_year(model: AdaptiveWeightModel, seq: dict, storage: str = "2day",
             e_terminal=None, pv_next: str = "recurse"):
    """全年滚动：返回 (x_plans, recs, targets)。

    storage : "daily"  旧口径日循环（计划/执行均 E0→E0，各日独立）；
              "2day"   2 日规划窗口 + 1 日执行（e_terminal=None 末端自由；
                       传入数值则锚定窗口末端）。
    """
    price = model.price
    x_plans = np.zeros((N_DAY, T))
    recs: list[dict | None] = [None] * N_DAY
    targets = np.zeros(N_DAY)
    E = E0
    for D in range(START_SIM, N_DAY):
        w, u = seq[D]
        l0, p0, l1, p1 = forecast_two_days(model, w, u, D, pv_next=pv_next)
        l0 = np.clip(l0 * KAPPA, 0.0, None) / 6.0
        l1 = np.clip(l1 * KAPPA, 0.0, None) / 6.0
        p0 = np.clip(p0 - MARGIN, 0.0, None) / 6.0
        p1 = np.clip(p1 - MARGIN, 0.0, None) / 6.0

        if storage == "daily":
            x_day, _E_day, _ = q2.plan_day(price, l0, p0, E0, eps=EPS_PLAN)
            e_end_target = E0
        else:
            x2, E2, _ = q2.plan_horizon(
                np.tile(price, 2), np.concatenate([l0, l1]),
                np.concatenate([p0, p1]), E, e_terminal, eps=EPS_PLAN)
            x_day = x2[:T]
            e_end_target = float(np.clip(E2[T - 1], E_MIN, E_MAX))

        ex = q2.exec_segment_causal(
            model.L[D] / 6.0, model.P[D] / 6.0, x_day, E, e_end_target)
        recs[D] = {
            "x": x_day,
            "c": ex["c"], "d": ex["d"], "s": ex["s"], "e": ex["e"], "E": ex["E"],
            "E_start": float(E), "E_end": float(ex["E"][-1]),
            "plan_cost": float(price @ x_day),
            "emerg_kwh": float(ex["e"].sum()),
            "emerg_cost": float(EMERG_MULT * (price @ ex["e"])),
            "spill_kwh": float(ex["s"].sum()),
        }
        x_plans[D] = x_day
        targets[D] = e_end_target
        E = float(ex["E"][-1]) if storage != "daily" else E0
    return x_plans, recs, targets


def summarize(recs, rep=REP) -> dict:
    """报送期汇总（万元/电量/储能轨迹统计）。"""
    plan = float(sum(recs[d]["plan_cost"] for d in rep))
    emerg = float(sum(recs[d]["emerg_cost"] for d in rep))
    return {
        "计划购电费/万元": round(plan / 1e4, 1),
        "紧急购电费/万元": round(emerg / 1e4, 1),
        "总费用/万元": round((plan + emerg) / 1e4, 1),
        "紧急购电量/kWh": float(sum(recs[d]["emerg_kwh"] for d in rep)),
        "发生紧急购电天数": int(sum(1 for d in rep if recs[d]["emerg_kwh"] > 1e-3)),
        "弃电量/kWh": float(sum(recs[d]["spill_kwh"] for d in rep)),
        "日末储电量最小值/kWh": round(min(recs[d]["E_end"] for d in rep), 1),
        "日末储电量最大值/kWh": round(max(recs[d]["E_end"] for d in rep), 1),
        "日末储电量均值/kWh": round(float(np.mean([recs[d]["E_end"] for d in rep])), 1),
        "日内储电量最小值/kWh": round(min(float(recs[d]["E"].min()) for d in rep), 1),
        "日内储电量最大值/kWh": round(max(float(recs[d]["E"].max()) for d in rep), 1),
        "_plan": plan, "_emerg": emerg,
    }


def verify(model: AdaptiveWeightModel, x_plans, recs, targets) -> dict:
    """约束回代：平衡、SOC 动态、区间、跨日衔接、执行目标跟踪。"""
    bal = dyn = 0.0
    emin, emax = np.inf, -np.inf
    track = 0.0
    for d in range(START_SIM, N_DAY):
        r = recs[d]
        res = (model.P[d] / 6.0 + r["x"] + r["d"] + r["e"]
               - model.L[d] / 6.0 - r["c"] - r["s"])
        bal = max(bal, float(np.abs(res).max()))
        Eser = np.concatenate([[r["E_start"]], r["E"]])
        dyn = max(dyn, float(np.abs(np.diff(Eser) - (ETA * r["c"] - r["d"] / ETA)).max()))
        emin = min(emin, float(r["E"].min()))
        emax = max(emax, float(r["E"].max()))
        track = max(track, abs(r["E_end"] - targets[d]))
    cont = max(abs(recs[d]["E_end"] - recs[d + 1]["E_start"])
               for d in range(START_SIM, N_DAY - 1))
    return {
        "功率平衡最大残差/kWh": bal,
        "储能动态最大残差/kWh": dyn,
        "储电量最小值/kWh": round(emin, 3),
        "储电量最大值/kWh": round(emax, 3),
        "跨日储电量衔接最大误差/kWh": cont,
        "执行末端目标跟踪最大误差/kWh": track,
    }


def verify_result2_file(path, model: AdaptiveWeightModel, x_plans, recs) -> dict:
    """回读 result2.xlsx：行数、勾稽、SOC 端点与跨日衔接。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    chk = {"sheet 名": "/".join(wb.sheetnames)}

    ws = wb["计划购电量"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    chk["计划表行数"] = len(rows)
    chk["全天购电量勾稽差/kWh"] = float(abs(
        sum(float(r[145]) for r in rows) - sum(float(x_plans[d].sum()) for d in REP)))
    chk["全天购电费勾稽差/元"] = float(abs(
        sum(float(r[146]) for r in rows) - sum(float(model.price @ x_plans[d]) for d in REP)))

    ws = wb["充放电量"]
    crows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] is not None]
    chk["充放表行数"] = len(crows)
    chk["充放表电量勾稽差/kWh"] = float(abs(
        sum(float(r[2]) for r in crows if r[2] is not None)
        - sum(float(recs[d]["c"].sum()) for d in REP)))
    soc_seq = [float(r[5]) for r in crows if r[4] in ("00:00", "24:00") and r[5] is not None]
    chk["储能端点最小值/kWh"] = round(min(soc_seq), 2)
    chk["储能端点最大值/kWh"] = round(max(soc_seq), 2)
    chk["跨日衔接最大误差/kWh"] = float(max(
        abs(soc_seq[2 * i + 1] - soc_seq[2 * i + 2]) for i in range(len(REP) - 1)))
    chk["表内0:00与记录差/kWh"] = float(max(
        abs(soc_seq[2 * i] - recs[d]["E_start"]) for i, d in enumerate(REP)))

    ws = wb["紧急购电量"]
    kw = sum(float(r[2]) for r in ws.iter_rows(min_row=2, values_only=True) if r[2] is not None)
    chk["紧急表电量勾稽差/kWh"] = float(abs(
        kw - sum(float(recs[d]["e"].sum()) for d in REP)))
    wb.close()
    return chk


# ---------------------------------------------------------------------------
# 蒙特卡洛：年度总成本分布（固定计划、重采样整日误差块）
# ---------------------------------------------------------------------------
def run_mc(model: AdaptiveWeightModel, x_plans, targets, e_start_feb: float,
           n_years: int = 100, seed: int = 42):
    """固定新结构计划，对 2.1-12.31 重采样附件 3 误差日块，评估紧急费用分布。"""
    data = model.data
    price = model.price
    Z, months, typ_hm, fc0 = data["Z"], data["months"], data["typ_hm"], data["fc0"]
    load = data["load"]
    rng = np.random.default_rng(seed)
    emerg_costs = np.zeros(n_years)
    emerg_kwhs = np.zeros(n_years)
    fallbacks = 0
    t0 = time.time()
    for y in range(n_years):
        E = float(e_start_feb)
        for d in REP:
            zi = int(rng.integers(0, N_DAY))
            pv_s = q2._hour_to_slots(
                q2._scenario_hourly(fc0[d], months[d], Z[zi], typ_hm)) / 6.0
            try:
                ex = q2.exec_segment_causal(
                    load[d] / 6.0, pv_s, x_plans[d], E, float(targets[d]))
            except RuntimeError:
                # 极端场景下规划末端不可达：退化为"维持当前储电量"（记数，不中断）
                fallbacks += 1
                ex = q2.exec_segment_causal(
                    load[d] / 6.0, pv_s, x_plans[d], E, E)
            emerg_kwhs[y] += float(ex["e"].sum())
            emerg_costs[y] += float(EMERG_MULT * (price @ ex["e"]))
            E = float(ex["E"][-1])
        if (y + 1) % 20 == 0:
            print(f"  MC {y + 1}/{n_years} 完成，用时 {time.time() - t0:.0f}s")
    return emerg_costs, emerg_kwhs, fallbacks


# ---------------------------------------------------------------------------
# 图表
# ---------------------------------------------------------------------------
def make_figures(model, recs, targets, daily, mc_costs=None, plan_cost=None) -> None:
    import matplotlib.pyplot as plt

    dates = model.dates

    fig, ax = pm.line(np.arange(len(REP)), [recs[d]["emerg_kwh"] for d in REP],
                      xlabel="日期", ylabel="紧急购电量 / kWh")
    month_starts = [i for i, d in enumerate(REP) if dates[d].endswith("-01")]
    ax.set_xticks(month_starts)
    ax.set_xticklabels([dates[REP[i]][5:7] + "月" for i in month_starts])
    pm.save_fig(fig, "Q2_逐日紧急购电",
                data=pd.DataFrame({"日期": [dates[d] for d in REP],
                                   "紧急购电量_kWh": [recs[d]["emerg_kwh"] for d in REP]}))

    # 储电量轨迹（日内最小/最大 + 日末规划目标）
    e_lo = np.array([float(recs[d]["E"].min()) for d in REP])
    e_hi = np.array([float(recs[d]["E"].max()) for d in REP])
    fig2, ax2 = pm.line(
        np.arange(len(REP)), [e_lo, e_hi, targets[REP]],
        labels=["日内最低", "日内最高", "日末规划目标"],
        xlabel="日期", ylabel="储电量 / kWh")
    ax2.axhline(E_MIN, color="grey", ls=":", lw=1)
    ax2.axhline(E_MAX, color="grey", ls=":", lw=1)
    ax2.set_xticks(month_starts)
    ax2.set_xticklabels([dates[REP[i]][5:7] + "月" for i in month_starts])
    pm.save_fig(fig2, "Q2_储能轨迹",
                data=pd.DataFrame({"日期": [dates[d] for d in REP],
                                   "日内最低/kWh": e_lo, "日内最高/kWh": e_hi,
                                   "日末目标/kWh": targets[REP]}))

    if mc_costs is not None and len(mc_costs) and plan_cost is not None:
        total = plan_cost + mc_costs
        fig3, ax3 = plt.subplots(figsize=(7, 4.3))
        ax3.hist(total, bins=30, color="#4C72B0", alpha=0.85, edgecolor="white")
        ax3.axvline(total.mean(), color="#C44E52", ls="--",
                    label=f"均值 {total.mean():,.0f} 元")
        ax3.axvline(np.percentile(total, 95), color="#55A868", ls=":",
                    label=f"P95 {np.percentile(total, 95):,.0f} 元")
        ax3.set_xlabel("年度总购电费 / 元")
        ax3.set_ylabel("频数")
        ax3.legend()
        pm.save_fig(fig3, "Q2_总费用分布",
                    data=pd.DataFrame({"年度总购电费_元": total}))


def probe(model: AdaptiveWeightModel, seq: dict) -> pd.DataFrame:
    """结构对比：日循环（旧口径）/ 2 日滚动·末端锚定 / 2 日滚动·末端自由。"""
    import matplotlib.pyplot as plt

    cfgs = [
        ("日循环（旧口径）", dict(storage="daily")),
        ("2 日滚动·末端锚定 6000", dict(storage="2day", e_terminal=E0)),
        ("2 日滚动·末端自由（新）", dict(storage="2day", e_terminal=None)),
    ]
    rows = []
    for name, kw in cfgs:
        t0 = time.time()
        x_plans, recs, targets = run_year(model, seq, **kw)
        s = summarize(recs)
        print(f"  [{name:<20}] 总 {s['总费用/万元']:8.1f} 万"
              f"（计划 {s['计划购电费/万元']:.1f} + 紧急 {s['紧急购电费/万元']:.1f}）"
              f" 紧急天数 {s['发生紧急购电天数']}，用时 {time.time() - t0:.0f}s")
        rows.append({
            "结构": name,
            "计划购电费/万元": s["计划购电费/万元"],
            "紧急购电费/万元": s["紧急购电费/万元"],
            "总费用/万元": s["总费用/万元"],
            "紧急购电量/kWh": round(s["紧急购电量/kWh"], 1),
            "紧急购电天数": s["发生紧急购电天数"],
            "日末储电量均值/kWh": s["日末储电量均值/kWh"],
            "日末储电量最小/kWh": s["日末储电量最小值/kWh"],
            "日末储电量最大/kWh": s["日末储电量最大值/kWh"],
        })
    df = pd.DataFrame(rows)
    pm.save_outputs(df, "q2_roll_structure_probe")

    fig, ax = plt.subplots(figsize=(7, 4.3))
    x = np.arange(len(df))
    ax.bar(x, df["计划购电费/万元"], label="计划购电费", color="#4C72B0")
    ax.bar(x, df["紧急购电费/万元"], bottom=df["计划购电费/万元"],
           label="紧急购电费", color="#C44E52")
    for i, v in enumerate(df["总费用/万元"]):
        ax.text(i, v + 15, f"{v:.1f}", ha="center", fontsize=8)
    ax.set_xticks(x)
    ax.set_xticklabels(df["结构"], fontsize=8)
    ax.set_ylabel("报送期费用 / 万元")
    ax.legend()
    pm.save_fig(fig, "Q2_结构对比", data=df)
    return df


# ---------------------------------------------------------------------------
# 官方 result2
# ---------------------------------------------------------------------------
def run_result2(model: AdaptiveWeightModel, seq: dict, mc_years: int = 0,
                seed: int = 42) -> dict:
    """生成官方 result2.xlsx（2 日滚动·末端自由）+ 校验 + 图表 + 报告。"""
    pm.init(seed=seed, root=str(ROOT))
    log = pm.get_logger("q2roll")
    t0 = time.time()
    x_plans, recs, targets = run_year(model, seq)
    s = summarize(recs)
    chk = verify(model, x_plans, recs, targets)
    print(f"新结构：总 {s['总费用/万元']:.1f} 万（计划 {s['计划购电费/万元']:.1f} + "
          f"紧急 {s['紧急购电费/万元']:.1f}），用时 {time.time() - t0:.0f}s")

    cur = RESULTS_DIR / "result2.xlsx"
    backup = pm.outputs_dir() / "result2_dailyCycle_tuned_backup.xlsx"
    if cur.exists() and not backup.exists():
        shutil.copy2(cur, backup)
        log.info("旧（日循环调优）result2 备份 -> {}", backup)

    out = q2.write_result2(model.dates, model.price, x_plans, recs, out_name="result2.xlsx")
    fchk = verify_result2_file(out, model, x_plans, recs)
    log.info("result2.xlsx -> {}", out)

    daily = pd.DataFrame({
        "日期": [model.dates[d] for d in REP],
        "计划购电量/kWh": [float(x_plans[d].sum()) for d in REP],
        "计划购电费/元": [recs[d]["plan_cost"] for d in REP],
        "紧急购电量/kWh": [recs[d]["emerg_kwh"] for d in REP],
        "紧急购电费/元": [recs[d]["emerg_cost"] for d in REP],
        "弃电量/kWh": [recs[d]["spill_kwh"] for d in REP],
        "日初储电量/kWh": [recs[d]["E_start"] for d in REP],
        "日末储电量/kWh": [recs[d]["E_end"] for d in REP],
        "日末规划目标/kWh": [targets[d] for d in REP],
    })
    pm.save_outputs(daily, "q2_roll_daily")

    mc_costs = None
    mc_note = ""
    if mc_years > 0:
        t1 = time.time()
        mc_costs, mc_kwhs, fallbacks = run_mc(
            model, x_plans, targets, recs[REPORT_START]["E_start"],
            n_years=mc_years, seed=seed)
        total = s["_plan"] + mc_costs
        p95 = float(np.percentile(total, 95))
        mc_stats = {
            "模拟年数": mc_years,
            "年度总成本均值/元": float(total.mean()),
            "年度总成本标准差/元": float(total.std()),
            "P5/元": float(np.percentile(total, 5)),
            "P50/元": float(np.percentile(total, 50)),
            "P95/元": p95,
            "CVaR95（尾部均值）/元": float(total[total >= p95].mean()),
            "年度紧急费用均值/元": float(mc_costs.mean()),
            "年度紧急费用P95/元": float(np.percentile(mc_costs, 95)),
            "末端目标不可达回退次数": fallbacks,
        }
        log.info("蒙特卡洛完成（{} 年，{:.0f}s）", mc_years, time.time() - t1)
        pm.save_outputs(
            pd.DataFrame({"计划费_元": s["_plan"], "紧急费_元": mc_costs,
                          "总费用_元": s["_plan"] + mc_costs}),
            "q2_roll_mc_annual_costs")
        mc_note = "年度分布：固定新结构计划，仅重采样附件 3 误差日块。"
        record("问题二 年度总成本分布（2 日滚动·蒙特卡洛）", mc_stats,
               note=("固定计划、场景重采样整日标准化误差块（保留日内相关）；"
                     "执行仍为逐槽因果（目标 = 各日规划末端储电量）；"
                     "图 figures/Q2_总费用分布.pdf。"))

    make_figures(model, recs, targets, daily, mc_costs, s["_plan"])

    record(
        "问题二 官方 result2（2 日滚动·末端自由）结果与校验",
        {
            "计划购电费/万元": s["计划购电费/万元"],
            "紧急购电费/万元": s["紧急购电费/万元"],
            "总费用/万元": s["总费用/万元"],
            "紧急购电量/kWh": round(s["紧急购电量/kWh"], 1),
            "发生紧急购电天数": s["发生紧急购电天数"],
            "弃电量/kWh": round(s["弃电量/kWh"], 1),
            "日末储电量最小值/kWh": s["日末储电量最小值/kWh"],
            "日末储电量最大值/kWh": s["日末储电量最大值/kWh"],
            "日末储电量均值/kWh": s["日末储电量均值/kWh"],
            "天数": len(REP),
            **chk, **fchk,
        },
        note=(
            "官方 results/result2.xlsx 由 2 日滚动（末端自由）结构生成：每天 0:00 以"
            "当日+次日 288 槽 LP 最小化计划购电费（末端不锚定 6000），只执行当日；"
            "预测为口径 E 调优配置（EWMA h=5 + κ=1.02 + m=50，次日无前视外推）；"
            "执行逐槽因果、缺口 5 倍紧急；储电量跨日连续（仿真自 2025-01-16 起，"
            "报送 2025-02-01 ~ 12-31）。旧日循环调优版备份 "
            "code/outputs/result2_dailyCycle_tuned_backup.xlsx。复现："
            "uv run python -m solve.q2_roll --result2 [--mc 100]。"
            f"{mc_note} 图 figures/Q2_逐日紧急购电.pdf、Q2_储能轨迹.pdf。"
        ),
    )
    return {"plan": s["_plan"], "emerg": s["_emerg"], "summary": s,
            "checks": {**chk, **fchk}, "out": out}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--result2", action="store_true", help="写官方 result2.xlsx")
    ap.add_argument("--probe", action="store_true", help="储能结构对比")
    ap.add_argument("--mc", type=int, default=0, help="蒙特卡洛年数（0 = 不做）")
    ap.add_argument("--pv-next", choices=["recurse", "same"], default="recurse",
                    help="窗口次日光伏预测口径")
    args = ap.parse_args()

    pm.init(seed=42, root=str(ROOT))
    model = AdaptiveWeightModel().load().build_table()
    seq = seqs_ewma_ext(model)
    print(f"权重序列就绪（EWMA h={HL:g}）：{len(seq)} 天；"
          f"仿真起点 2025-01-{START_SIM + 1:02d}，报送 2025-02-01 ~ 12-31")

    if args.probe:
        df = probe(model, seq)
        print(df.to_string(index=False))
        record("问题二 储能结构对比（日循环 vs 2 日滚动）", df,
               note=("新口径为 2 日滚动·末端自由（不锚定 6000）：每天 2 日规划窗口、"
                     "只执行当日、储电量跨日连续；另列旧日循环（当日→当日）与"
                     "2 日滚动·末端锚定 6000 作对照。计划/执行其余口径一致"
                     "（EWMA h=5 + κ=1.02 + m=50；逐槽因果执行）。"
                     "图 figures/Q2_结构对比.pdf。"))
        return
    if args.result2:
        run_result2(model, seq, mc_years=args.mc, seed=42)
        return
    ap.print_help()


if __name__ == "__main__":
    main()
