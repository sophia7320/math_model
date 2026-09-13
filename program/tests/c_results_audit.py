"""C 题五份官方结果文件的结构审计（独立于写出脚本，提交前固定检查）。

校验内容：sheet 名与顺序、日期连续性、行数、数值非负/有限、全天购电量与
源电价逐槽费用重算、储能端点/跨日连续性与 SOC 界、时间标签。

注意：官方模板（附件 5）把 24:00 储电量放在 4:00-8:00 行，本脚本按模板校验。

运行（program/ 下）：uv run python tests/c_results_audit.py
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "results"
DATA_C = ROOT / "program" / "data" / "C"

EXPECTED = {
    "result1.xlsx": (["计划购电量", "充放电量"], 1, 6),
    "result2.xlsx": (["计划购电量", "充放电量", "紧急购电量"], 334, 2004),
    "result3.xlsx": (["计划购电量", "调整购电量", "充放电量", "紧急购电量"], 334, 2004),
    "result4-2.xlsx": (["计划购电量", "充放电量", "紧急购电量"], 334, 2004),
    "result4-3.xlsx": (["计划购电量", "调整购电量", "充放电量", "紧急购电量"], 334, 2004),
}


def _as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)).date()


def _prices() -> tuple[np.ndarray, dict[date, np.ndarray]]:
    wb = openpyxl.load_workbook(DATA_C / "附件1.xlsx", read_only=True, data_only=True)
    static = np.asarray([r[1] for r in list(wb.active.values)[1:145]], dtype=float)
    wb.close()
    wb = openpyxl.load_workbook(DATA_C / "附件4.xlsx", read_only=True, data_only=True)
    dynamic = {
        _as_date(r[0]): np.asarray(r[1:145], dtype=float)
        for r in list(wb.active.values)[1:] if r[0] is not None
    }
    wb.close()
    return static, dynamic


def audit_workbook(name: str, sheets: list[str], n_days: int, n_charge: int) -> None:
    path = RESULTS / name
    assert path.exists(), f"缺少 {path}"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    assert wb.sheetnames == sheets, (name, wb.sheetnames)

    # 计划/调整购电量：日行 × 144 槽 + 全天电量/费用汇总列
    if name == "result1.xlsx":
        rows = [r for r in wb["计划购电量"].iter_rows(min_row=2, values_only=True)
                if r[0] is not None]
        assert len(rows) == 144, (name, "计划购电量", len(rows))
        values = np.asarray([r[1] for r in rows], dtype=float)
        assert np.isfinite(values).all() and (values >= -1e-7).all()
    static_price, dynamic_price = _prices()
    purchase_rows: dict[str, list[tuple]] = {}
    for sheet in ("计划购电量", "调整购电量"):
        if sheet not in sheets or name == "result1.xlsx":
            continue
        rows = [r for r in wb[sheet].iter_rows(min_row=2, values_only=True) if r[0] is not None]
        purchase_rows[sheet] = rows
        assert len(rows) == n_days, (name, sheet, len(rows))
        dates = [_as_date(r[0]) for r in rows]
        expected = [date(2025, 2, 1) + timedelta(days=i) for i in range(n_days)]
        assert dates == expected, (name, sheet, "日期不连续或范围错误", dates[:1], dates[-1:])
        for i, r in enumerate(rows):
            slots = np.asarray(r[1:145], dtype=float)
            assert np.isfinite(slots).all() and (slots >= -1e-7).all()
            assert abs(float(r[145]) - float(slots.sum())) < 1e-4, (name, sheet, "勾稽")
            assert np.isfinite(float(r[146])) and float(r[146]) >= -1e-7
            price = dynamic_price[dates[i]] if name.startswith("result4-") else static_price
            if sheet == "调整购电量":
                plan = np.asarray(purchase_rows["计划购电量"][i][1:145], dtype=float)
                expected_fee = float(price @ slots + (0.5 * price * np.abs(slots - plan)).sum())
            else:
                expected_fee = float(price @ slots)
            assert abs(float(r[146]) - expected_fee) < 1e-4, (
                name, sheet, "费用列不能由源电价重算", float(r[146]), expected_fee)

    # 充放电量：0:00/24:00 行（官方模板位于第 1/2 块）标签与 SOC
    charge = [r for r in wb["充放电量"].iter_rows(min_row=2, values_only=True)
              if r[1] is not None]
    assert len(charge) == n_charge, (name, "充放电量", len(charge))
    cc, dc, ec = (1, 2, 4) if name == "result1.xlsx" else (2, 3, 5)
    e0 = e24 = None
    first_block = charge[0]
    second_block = charge[1] if len(charge) > 1 else None
    if name == "result1.xlsx":
        assert first_block[3] == "0:00" and first_block[ec] is not None
        assert second_block[3] == "24:00" and second_block[ec] is not None
        e0, e24 = float(first_block[ec]), float(second_block[ec])
    for r in charge:
        assert float(r[cc]) >= -1e-7 and float(r[dc]) >= -1e-7
        if r[ec] is not None:
            assert 1200.0 - 1e-6 <= float(r[ec]) <= 10800.0 + 1e-6, (name, r[ec])
    if name == "result1.xlsx":
        assert abs(e0 - e24) < 1e-6, (name, "0:00/24:00 储电量不相等", e0, e24)
    else:
        plan_dates = [_as_date(r[0]) for r in purchase_rows["计划购电量"]]
        e0s, e24s, charge_dates = [], [], []
        for i in range(n_days):
            block = charge[6 * i:6 * (i + 1)]
            assert block[0][4] == "00:00" and block[1][4] == "24:00", (
                name, i, "0:00/24:00 标签错误")
            charge_dates.append(_as_date(block[0][0]))
            e0s.append(float(block[0][ec]))
            e24s.append(float(block[1][ec]))
        assert charge_dates == plan_dates, (name, "充放表日期与计划表不一致")
        assert max(abs(e24s[i] - e0s[i + 1]) for i in range(n_days - 1)) < 1e-6, (
            name, "跨日 SOC 不连续")

    # 紧急购电量：电量非负
    if "紧急购电量" in sheets:
        for r in wb["紧急购电量"].iter_rows(min_row=2, values_only=True):
            if r[2] is not None:
                assert float(r[2]) >= -1e-7, (name, r[:3])

    # 时间标签：与附件 5 模板一致（附件与模板标签相差一个 10 分钟刻度，
    # 一律按行序号对齐；重写文件时不得改动模板标签）
    tpl = openpyxl.load_workbook(DATA_C / "附件5" / name, read_only=True, data_only=True)
    t_rows = list(tpl["计划购电量"].iter_rows(values_only=True))
    hdr = next(wb["计划购电量"].iter_rows(min_row=1, max_row=1, values_only=True))

    def _norm(v):
        return None if v is None else str(v)

    assert [_norm(v) for v in hdr] == [_norm(v) for v in t_rows[0]], (
        name, "计划购电量表头/时段标签与附件 5 模板不一致")
    if name == "result1.xlsx":
        labels = [_norm(r[0]) for r in wb["计划购电量"].iter_rows(min_row=2, values_only=True)
                  if r[0] is not None]
        assert labels == [_norm(r[0]) for r in t_rows[1:]], (
            name, "逐行时段标签与附件 5 模板不一致")
    label_col = 0 if name == "result1.xlsx" else 1   # 充放电量表的“时间段”列
    blocks = [str(r[label_col]) for r in wb["充放电量"].iter_rows(min_row=2, values_only=True)
              if r[label_col] is not None]
    expected_blocks = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
                       "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    assert blocks and all(b == expected_blocks[i % 6] for i, b in enumerate(blocks)), (
        name, "充放电量块标签错误", blocks[:6])
    tpl.close()
    wb.close()


def main() -> None:
    for name, (sheets, n_days, n_charge) in EXPECTED.items():
        audit_workbook(name, sheets, n_days, n_charge)
        print(f"ok - {name}")
    print("C result audit: PASS")


if __name__ == "__main__":
    main()
